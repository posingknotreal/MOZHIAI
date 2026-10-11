import io
import base64
import http.client
import json
import os
import sys
import tempfile
import threading
import unittest
from urllib.parse import urlsplit
import urllib.error
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app


def make_pdf():
    from reportlab.pdfgen import canvas
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 760, "MOZHIAI test file. This page discusses classical poetry and grammar.")
    pdf.showPage()
    pdf.drawString(72, 760, "Gradient descent updates model parameters to reduce the loss function during training.")
    pdf.drawString(72, 740, "The learning rate controls the size of each update step.")
    pdf.save()
    return output.getvalue()


def make_docx():
    from docx import Document
    output = io.BytesIO()
    doc = Document()
    doc.add_heading("Tamil research notes", 0)
    doc.add_paragraph("Tamil Unicode text: வணக்கம். Machine learning models learn patterns from examples.")
    doc.save(output)
    return output.getvalue()


def multipart(files):
    boundary = "mozhi-test-boundary-987"
    body = bytearray()
    for name, content_type, data in files:
        body.extend(f"--{boundary}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"{name}\"\r\nContent-Type: {content_type}\r\n\r\n".encode())
        body.extend(data)
        body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode())
    return boundary, bytes(body)


def request_bytes(base, path, method="GET", body=None, headers=None):
    target = urlsplit(base + path)
    connection = http.client.HTTPConnection(target.hostname, target.port, timeout=10)
    connection.request(method, target.path + ("?" + target.query if target.query else ""), body=body, headers=headers or {})
    response = connection.getresponse()
    content = response.read()
    status, response_headers = response.status, response.headers
    connection.close()
    if status >= 400:
        raise urllib.error.HTTPError(base + path, status, "HTTP request failed", response_headers, io.BytesIO(content))
    return content


class AppChecks(unittest.TestCase):
    def test_deployment_auth_protects_static_and_api_routes(self):
        with patch.dict(os.environ, {"MOZHI_AUTH_USER": "owner", "MOZHI_AUTH_PASSWORD": "a-long-test-password"}, clear=True):
            server = ThreadingHTTPServer(("127.0.0.1", 0), app.Handler)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                for path in ("/", "/api/documents"):
                    with self.assertRaises(urllib.error.HTTPError) as caught:
                        request_bytes(base, path)
                    self.assertEqual(caught.exception.code, 401)
                token = base64.b64encode(b"owner:a-long-test-password").decode()
                html = request_bytes(base, "/", headers={"Authorization": "Basic " + token})
                self.assertIn(b"MOZHIAI", html)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_normalization_and_tanglish_expansion(self):
        self.assertEqual(app.normalize("வணக்கம்\u200b   world"), "வணக்கம் world")
        self.assertIn("இந்த", app.tokenize("indha document la enna irukku?"))
        self.assertIn("என்ன", app.tokenize("indha document la enna irukku?"))
        self.assertIn("into English", app.translation_request("வணக்கம்", ""))
        self.assertIn("into Tamil", app.translation_request("Hello there", "Translate this into Tamil"))

    def test_vectors_and_page_preserving_chunks(self):
        vector = app.local_vector("தமிழ் document retrieval")
        self.assertAlmostEqual(sum(x*x for x in vector), 1.0, places=6)
        parts = app.split_chunks([(7, " ".join(["மொழி"] * 500))], max_words=100, overlap=20)
        self.assertTrue(parts)
        self.assertTrue(all(page == 7 for page, _ in parts))
        self.assertGreater(len(parts), 1)

    def test_server_side_provider_adapters(self):
        payloads = []
        def fake_http(url, payload, key=None, timeout=90):
            payloads.append((url, payload, key))
            if url.endswith("/embeddings"):
                values = payload["input"] if isinstance(payload["input"], list) else [payload["input"]]
                return {"data": [{"index": i, "embedding": [float(i + 1), 0.5]} for i, _ in enumerate(values)]}
            return {"choices": [{"message": {"content": "Grounded Tamil answer."}}]}

        with patch.dict(os.environ, {"MOZHI_LLM_BASE_URL": "http://provider.local/v1", "MOZHI_LLM_API_KEY": "server-secret",
                                     "MOZHI_LLM_MODEL": "test-chat", "MOZHI_EMBEDDING_BASE_URL": "http://embed.local/v1",
                                     "MOZHI_EMBEDDING_API_KEY": "embed-secret", "MOZHI_EMBEDDING_MODEL": "test-embed"}, clear=True), \
                patch.object(app, "http_json", side_effect=fake_http):
            vectors = app.embed_many(["தமிழ்", "Tanglish"])
            self.assertEqual(vectors, [[1.0, 0.5], [2.0, 0.5]])
            self.assertEqual(app.embed_query("கேள்வி"), [1.0, 0.5])
            answer = app.model_complete([{"role": "user", "content": "hello"}])
            self.assertEqual(answer, "Grounded Tamil answer.")
            self.assertTrue(app.provider_config()["llm"])
            self.assertTrue(app.provider_config()["embedding"])
            self.assertNotIn("server-secret", json.dumps(app.provider_config()))
            self.assertNotIn("embed-secret", json.dumps(app.provider_config()))
            self.assertEqual(payloads[-1][2], "server-secret")

    def test_optional_qdrant_adapter_and_document_filter(self):
        calls = []
        def fake_request(method, path, payload=None, timeout=25):
            calls.append((method, path, payload))
            return {"result": [{"payload": {"chunk_id": "abc"}, "score": 0.82}]}
        with patch.dict(os.environ, {"MOZHI_VECTOR_BACKEND": "qdrant", "MOZHI_QDRANT_URL": "http://qdrant.local",
                                     "MOZHI_QDRANT_API_KEY": "vector-secret"}, clear=True), \
                patch.object(app, "qdrant_collection", return_value="mozhiai"), \
                patch.object(app, "qdrant_request", side_effect=fake_request):
            scores = app.query_qdrant([0.1, 0.2], ["document-1"], 5)
            self.assertEqual(scores, {"abc": 0.82})
            self.assertEqual(calls[0][0], "POST")
            self.assertIn("points/search", calls[0][1])
            self.assertEqual(calls[0][2]["filter"]["should"][0]["match"]["value"], "document-1")
            self.assertNotIn("vector-secret", json.dumps(app.provider_config()))

    def test_document_extractors(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            txt = root / "tamil.txt"
            txt.write_text("வணக்கம்\nTamil and English", encoding="utf-8")
            pages, count = app.extract_doc(txt, ".txt")
            self.assertIsNone(count)
            self.assertIn("வணக்கம்", pages[0][1])
            pdf = root / "two.pdf"
            pdf.write_bytes(make_pdf())
            pages, count = app.extract_doc(pdf, ".pdf")
            self.assertEqual(count, 2)
            self.assertEqual([page for page, _ in pages], [1, 2])
            self.assertIn("Gradient descent", pages[1][1])
            docx = root / "notes.docx"
            docx.write_bytes(make_docx())
            pages, count = app.extract_doc(docx, ".docx")
            self.assertIsNone(count)
            self.assertIn("வணக்கம்", pages[0][1])

    def test_optional_ocr_provider_for_scanned_pdf(self):
        from reportlab.pdfgen import canvas
        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        pdf.showPage()
        pdf.save()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "scan.pdf"
            path.write_bytes(output.getvalue())
            class FakeResponse:
                def __enter__(self): return self
                def __exit__(self, *args): return None
                def read(self): return json.dumps({"pages": [{"page": 1, "text": "OCR text: வணக்கம்"}]}).encode()
            seen = {}
            def fake_urlopen(request, timeout=90):
                seen["url"] = request.full_url
                seen["method"] = request.get_method()
                seen["body"] = request.data
                seen["headers"] = {key.casefold(): value for key, value in request.header_items()}
                return FakeResponse()
            with patch.dict(os.environ, {"MOZHI_OCR_BASE_URL": "http://ocr.local/api", "MOZHI_OCR_API_KEY": "ocr-secret"}, clear=True), \
                    patch.object(app.urllib.request, "urlopen", side_effect=fake_urlopen):
                pages, count = app.extract_doc(path, ".pdf")
            self.assertEqual(count, 1)
            self.assertIn("OCR text", pages[0][1])
            self.assertEqual(seen["url"], "http://ocr.local/api/ocr")
            self.assertEqual(seen["method"], "POST")
            self.assertEqual(seen["body"], path.read_bytes())
            self.assertEqual(seen["headers"]["authorization"], "Bearer ocr-secret")

    def test_local_tesseract_fallback_only_processes_blank_pdf_pages(self):
        from reportlab.pdfgen import canvas
        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        pdf.showPage()
        pdf.drawString(72, 760, "This page already contains extractable text.")
        pdf.save()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "mixed.pdf"
            path.write_bytes(output.getvalue())
            with patch.dict(os.environ, {}, clear=True), \
                    patch.object(app, "tesseract_ocr_pdf", return_value=[(1, "OCR text: வணக்கம்")]) as local_ocr:
                pages, count = app.extract_doc(path, ".pdf")
            self.assertEqual(count, 2)
            self.assertIn("OCR text", pages[0][1])
            self.assertIn("extractable text", pages[1][1])
            local_ocr.assert_called_once_with(path, [1])

    def test_ocr_status_reports_local_tesseract_without_cloud_provider(self):
        with patch.dict(os.environ, {}, clear=True), patch.object(app, "tesseract_command", return_value="tesseract"):
            status = app.provider_config()
        self.assertTrue(status["ocr"])
        self.assertEqual(status["ocrMode"], "local")

    def test_http_upload_retrieval_chat_history_and_translation_fallback(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {}, clear=True):
            root = Path(temp)
            with patch.object(app, "DATA", root), patch.object(app, "FILES", root / "files"), patch.object(app, "DB", root / "test.sqlite3"):
                app.FILES.mkdir()
                app.init_db()
                server = ThreadingHTTPServer(("127.0.0.1", 0), app.Handler)
                thread = threading.Thread(target=server.serve_forever, daemon=True)
                thread.start()
                base = f"http://127.0.0.1:{server.server_port}"
                try:
                    boundary, body = multipart([
                        ("research.pdf", "application/pdf", make_pdf()),
                        ("notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx()),
                        ("plain.txt", "text/plain", "Tamil document storage works.".encode()),
                    ])
                    request = urllib.request.Request(base + "/api/documents", data=body, method="POST",
                                                     headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
                    uploaded = json.loads(request_bytes(base, "/api/documents", "POST", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"}))
                    self.assertEqual(len(uploaded["documents"]), 3)
                    for _ in range(50):
                        current = json.loads(request_bytes(base, "/api/metrics"))["documents"]
                        if all(d["status"] not in {"processing", "extracting", "indexing"} for d in current):
                            break
                        threading.Event().wait(0.05)
                    uploaded["documents"] = current
                    self.assertTrue(all(d["status"] == "ready" for d in uploaded["documents"]))
                    pdf_doc = next(d for d in uploaded["documents"] if d["extension"] == ".pdf")
                    search = json.loads(request_bytes(base, "/api/search?q=gradient%20descent"))
                    self.assertEqual(search["results"][0]["page"], 2)
                    chat_request = urllib.request.Request(base + "/api/chat", data=json.dumps({"message": "gradient descent", "documentIds": [pdf_doc["id"]]}).encode(), method="POST", headers={"Content-Type":"application/json"})
                    chat = json.loads(request_bytes(base, "/api/chat", "POST", json.dumps({"message": "gradient descent", "documentIds": [pdf_doc["id"]]}).encode(), {"Content-Type":"application/json"}))
                    self.assertEqual(chat["citations"][0]["page"], 2)
                    follow_request = urllib.request.Request(base + "/api/chat", data=json.dumps({"chatId": chat["chatId"], "message": "What sets the update size?", "documentIds": [pdf_doc["id"]]}).encode(), method="POST", headers={"Content-Type":"application/json"})
                    request_bytes(base, "/api/chat", "POST", json.dumps({"chatId": chat["chatId"], "message": "What sets the update size?", "documentIds": [pdf_doc["id"]]}).encode(), {"Content-Type":"application/json"})
                    messages = json.loads(request_bytes(base, "/api/chats/" + chat["chatId"]))["messages"]
                    self.assertEqual(len(messages), 4)
                    action_request = urllib.request.Request(base + "/api/action", data=json.dumps({"mode": "translate", "text": "வணக்கம்"}).encode(), method="POST", headers={"Content-Type":"application/json"})
                    try:
                        request_bytes(base, "/api/action", "POST", json.dumps({"mode": "translate", "text": "வணக்கம்"}).encode(), {"Content-Type":"application/json"})
                        self.fail("Offline translation should clearly require a model.")
                    except urllib.error.HTTPError as exc:
                        self.assertEqual(exc.code, 503)
                finally:
                    server.shutdown()
                    server.server_close()
                    thread.join(timeout=2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
